"""Excel çıktıları: şablon dosyalarının (templates/) üzerine yazar, biçimi korur."""
import calendar
import io
import os
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

import model

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates')
WHITE = PatternFill(fill_type='solid', start_color='FFFFFF', end_color='FFFFFF')


def _fill(hex_color):
    c = (hex_color or '#FFFFFF').lstrip('#').upper()
    return PatternFill(fill_type='solid', start_color=c, end_color=c)


def _load(name):
    return openpyxl.load_workbook(os.path.join(TEMPLATE_DIR, name))


def _save(wb):
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def _shift_rows(ws, start, amount):
    """amount > 0: `start` satırından önce satır ekler; amount < 0: `start`tan itibaren satır siler.
    openpyxl birleştirilmiş hücreleri, satır yüksekliklerini ve yazdırma alanını kaydırmadığı için
    bunları burada elle kaydırıyoruz."""
    if amount == 0:
        return
    end_deleted = start - amount - 1 if amount < 0 else None
    merged = []
    for rng in list(ws.merged_cells.ranges):
        bounds = (rng.min_col, rng.min_row, rng.max_col, rng.max_row)
        ws.unmerge_cells(str(rng))
        merged.append(bounds)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height is not None}

    if amount > 0:
        ws.insert_rows(start, amount)
    else:
        ws.delete_rows(start, -amount)

    def move(r):
        return r + amount if r >= start else r

    for c1, r1, c2, r2 in merged:
        if amount < 0 and r1 >= start and r2 <= end_deleted:
            continue  # tamamen silinen alan
        ws.merge_cells(start_row=move(r1), start_column=c1, end_row=move(r2), end_column=c2)

    for r in list(ws.row_dimensions.keys()):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        if amount < 0 and start <= r <= end_deleted:
            continue
        ws.row_dimensions[move(r)].height = h

    if ws.print_area:
        area = ws.print_area.split('!')[-1].replace('$', '')
        first, last = area.split(':')
        col = ''.join(ch for ch in last if ch.isalpha())
        row = int(''.join(ch for ch in last if ch.isdigit()))
        ws.print_area = '%s:%s%d' % (first, col, move(row))


def _month_context(doc, ym):
    doc = model.normalize(doc)
    year, mon = model.parse_ym(ym)
    return doc, year, mon, model.days_in_month(ym), model.get_month(doc, ym)


def _title(year, mon):
    return '%d %s AYI' % (year, model.MONTHS_TR[mon - 1])


def build_puantaj(doc, ym):
    doc, year, mon, dim, month = _month_context(doc, ym)
    settings = doc['settings']
    kurum, imza = settings['kurum'], settings['imza']
    excuses = model.excuse_map(settings)
    physicians = model.puantaj_order(model.month_physicians(doc, ym), settings)
    holiday_fill = _fill(settings['colors']['holiday'])

    wb = _load('puantaj.xlsx')
    ws = wb.worksheets[0]
    ws.title = model.MONTHS_TR[mon - 1]

    ws['A1'].value = ' T.C.\n%s VALİLİĞİ\n%s\n%s\n%s  BİLDİRİMİ PUANTAJ ÇİZELGESİ' % (
        model.tr_upper(kurum['il']), kurum['mudurluk'], kurum['hastane'], _title(year, mon))
    ws['A2'].value = 'Birim Adı: %s' % kurum['birim']

    # İmza alanları (şablondaki sabit konumlar; satır kaydırmadan önce yazılır)
    ws['M30'].value = imza['hazirlayanBaslik']
    ws['J33'].value = imza['hazirlayanUnvan']
    ws['AJ30'].value = imza['onayBaslik']
    ws['AJ33'].value = imza['onayUnvan']
    for cell, merge, name in (('J32', 'J32:P32', imza['hazirlayanAd']),
                              ('AJ32', 'AJ32:AN32', imza['onayAd'])):
        if name:
            ws.merge_cells(merge)
            ws[cell].value = name
            ws[cell].alignment = Alignment(horizontal='center')

    # Gün başlıkları (satır 3): ay kısa ise fazla günleri boşalt, tatilleri gölgele
    for day in range(1, 32):
        cell = ws.cell(row=3, column=4 + day)
        cell.value = day if day <= dim else None
        cell.fill = holiday_fill if day <= dim and day in month['holidays'] else WHITE

    # Hekim satırları: şablonda 4-21 arası 18 satır var
    first, slots = 4, 18
    needed = max(len(physicians), 1)
    style_row = [ws.cell(row=first, column=c) for c in range(1, ws.max_column + 1)]
    if needed > slots:
        _shift_rows(ws, first + slots, needed - slots)
        for r in range(first + slots, first + needed):
            for src in style_row:
                _copy_style(src, ws.cell(row=r, column=src.column))
            ws.row_dimensions[r].height = ws.row_dimensions[first].height
    elif needed < slots:
        _shift_rows(ws, first + needed, needed - slots)

    for i, p in enumerate(physicians):
        row = first + i
        ws.cell(row=row, column=1).value = i + 1
        ws.cell(row=row, column=2).value = p.get('tc') or None
        ws.cell(row=row, column=3).value = model.tr_upper(p['name'])
        ws.cell(row=row, column=4).value = 'UZM.DR.' if model.is_specialist(p, settings) else 'TABİP'
        for day in range(1, 32):
            val = month['schedule'].get('%s-%d' % (p['id'], day), '') if day <= dim else ''
            if val in model.SHIFT_HOURS:
                out = int(val)
            elif val in excuses:
                out = excuses[val].get('puantaj') or None
            else:
                out = None
            ws.cell(row=row, column=4 + day).value = out
    if not physicians:
        for c in range(1, 36):
            ws.cell(row=first, column=c).value = None

    return _save(wb), 'Puantaj_%s_%d.xlsx' % (model.MONTHS_TR[mon - 1], year)


def _date_rows(ws, first_row, year, mon, dim, holidays, holiday_fill, clear_cols):
    for day in range(1, 32):
        row = first_row + day - 1
        a, b = ws.cell(row=row, column=1), ws.cell(row=row, column=2)
        if day <= dim:
            a.value = datetime(year, mon, day)
            a.number_format = 'DD.MM.YYYY'
            b.value = model.DAYS_TR[calendar.weekday(year, mon, day)]
            b.fill = holiday_fill if day in holidays else WHITE
        else:
            a.value = b.value = None
            b.fill = WHITE
        for col in clear_cols:
            ws.cell(row=row, column=col).value = None


def build_isimler(doc, ym):
    doc, year, mon, dim, month = _month_context(doc, ym)
    settings = doc['settings']
    kurum = settings['kurum']
    rows, doctors = model.isimler_rows(doc, ym)

    wb = _load('nobetisimler_template.xlsx')
    ws = wb['Nöbet listesi']
    ws['A1'].value = 'T.C.'
    ws['A2'].value = '%s VALİLİĞİ' % model.tr_upper(kurum['il'])
    ws['A3'].value = model.tr_upper(kurum['mudurluk'])
    ws['A4'].value = model.tr_upper(kurum['hastane'])
    ws['A5'].value = '  %s NÖBET LİSTESİ' % _title(year, mon)

    first = 7
    _date_rows(ws, first, year, mon, dim, month['holidays'], _fill(settings['colors']['holiday']),
               clear_cols=range(3, 8))
    for r in rows:
        row = first + r['day'] - 1
        ws.cell(row=row, column=3).value = r['mesai'] or None
        for i, name in enumerate(r['nobet']):
            ws.cell(row=row, column=4 + i).value = name or None
    for i, name in enumerate(doctors):
        ws.cell(row=first + i, column=7).value = name

    return _save(wb), 'Nobet_Isimleri_%s_%d.xlsx' % (model.MONTHS_TR[mon - 1], year)


def build_renkli(doc, ym):
    doc, year, mon, dim, month = _month_context(doc, ym)
    settings = doc['settings']
    excuses = model.excuse_map(settings)
    physicians = model.month_physicians(doc, ym)
    excuse_fill = _fill(settings['colors']['excuseExcel'])

    wb = _load('nobetprogramirenkli.xlsx')
    ws = wb['Nöbet listesi']
    header, first, base_cols = 9, 10, 13   # şablonda C..O = 13 hekim sütunu

    # Şablondan fazla hekim varsa son sütunun biçimini kopyala
    last_col = 2 + base_cols
    for i in range(base_cols, len(physicians)):
        col = 3 + i
        for row in range(header, first + 31):
            _copy_style(ws.cell(row=row, column=last_col), ws.cell(row=row, column=col))
        ws.column_dimensions[get_column_letter(col)].width = \
            ws.column_dimensions[get_column_letter(last_col)].width or 11

    used_cols = range(3, 3 + max(base_cols, len(physicians)))
    _date_rows(ws, first, year, mon, dim, month['holidays'], _fill(settings['colors']['holiday']),
               clear_cols=used_cols)
    for col in used_cols:
        ws.cell(row=header, column=col).value = None
        ws.cell(row=header, column=col).fill = WHITE
        for day in range(1, 32):
            ws.cell(row=first + day - 1, column=col).fill = WHITE

    for i, p in enumerate(physicians):
        col = 3 + i
        head = ws.cell(row=header, column=col)
        head.value = p.get('displayName') or p['name']
        head.fill = _fill(p['color'])
        for day in range(1, dim + 1):
            cell = ws.cell(row=first + day - 1, column=col)
            val = month['schedule'].get('%s-%d' % (p['id'], day), '')
            if val == '24':
                cell.fill = _fill(p['color'])
            elif val in ('16', '8'):
                cell.fill = _fill(p['color'])
                cell.value = int(val)
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif val in excuses:
                cell.fill = excuse_fill

    return _save(wb), 'Nobet_Listesi_%s_%d.xlsx' % (model.MONTHS_TR[mon - 1], year)


BUILDERS = {'puantaj': build_puantaj, 'isimler': build_isimler, 'renkli': build_renkli}
